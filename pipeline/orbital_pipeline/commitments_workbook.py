"""Read one Finance commitments workbook into clean contract rows.

The layout is not stable between exports, so nothing here depends on positions:
- the sheet is the one that contains a recognisable header row (other sheets are ignored);
- the header row is found by scanning the top of the sheet (title rows above it are ignored);
- columns are matched by name, using a small list of known aliases (order doesn't matter,
  unknown columns are ignored);
- rows that aren't contracts (blank, TOTAL, footnotes) are skipped.

When the layout changes in a way we don't recognise, this module raises WorkbookLayoutError
with the headers it did find. The run fails loudly instead of loading wrong numbers.
"""

import logging
import re
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from openpyxl.utils.datetime import from_excel

log = logging.getLogger(__name__)

SNAPSHOT_FILE_PATTERN = re.compile(r"^commitments_(\d{4}-\d{2}-\d{2})\.xlsx$")
CONTRACT_ID_PATTERN = re.compile(r"^K-\d+$")
HEADER_SEARCH_ROWS = 30

# Known spellings of each column, after normalising (lowercase, letters and digits only).
# To support a new spelling, add it here.
COLUMN_ALIASES = {
    "contract_id":    {"contractid", "contract", "contractno", "contractnumber"},
    "customer_id":    {"customerid", "custid", "customerno", "customernumber"},
    "customer_name":  {"customer", "customername", "account", "accountname"},
    "start_date":     {"startdate", "start", "termstart", "contractstart"},
    "end_date":       {"enddate", "end", "termend", "contractend"},
    "commitment_usd": {"commitusd", "commitmentusd", "totalcommitmentusd", "commitment",
                       "commit", "totalcommitment", "commitmentamountusd"},
    "status":         {"status", "contractstatus"},
}
REQUIRED_COLUMNS = ("contract_id", "customer_id", "start_date", "end_date", "commitment_usd")

DATE_FORMATS = ("%Y-%m-%d", "%d %b %Y", "%d %B %Y", "%d/%m/%Y", "%d.%m.%Y", "%b %d, %Y")


class WorkbookLayoutError(ValueError):
    pass


@dataclass(frozen=True)
class ContractTerms:
    """The contract attributes we track over time (the SCD2 'tracked columns')."""
    customer_id: str
    customer_name: str | None
    start_date: date
    end_date: date
    commitment_usd: Decimal
    status: str | None


@dataclass(frozen=True)
class SnapshotRow:
    contract_id: str
    terms: ContractTerms
    source_sheet: str
    source_row: int


def snapshot_date_from_filename(path: Path) -> date | None:
    match = SNAPSHOT_FILE_PATTERN.match(path.name)
    return date.fromisoformat(match.group(1)) if match else None


def read_workbook(path: Path) -> list[SnapshotRow]:
    """Return one row per contract in the workbook. Exact duplicate rows are dropped."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet, header_row_number, columns = _find_contract_table(workbook, path.name)
        rows_by_contract: dict[str, SnapshotRow] = {}
        for row_number, cells in enumerate(
                sheet.iter_rows(min_row=header_row_number + 1, values_only=True),
                start=header_row_number + 1):
            row = _parse_row(cells, columns, sheet.title, row_number, path.name)
            if row is None:
                continue
            existing = rows_by_contract.get(row.contract_id)
            if existing is None:
                rows_by_contract[row.contract_id] = row
            elif existing.terms == row.terms:
                log.warning("%s: contract %s appears twice with identical terms (rows %d and %d); "
                            "keeping one", path.name, row.contract_id, existing.source_row, row_number)
            else:
                raise ValueError(f"{path.name}: contract {row.contract_id} appears twice with "
                                 f"different terms (rows {existing.source_row} and {row_number})")
        if not rows_by_contract:
            raise WorkbookLayoutError(f"{path.name}: found the header but no contract rows")
        return list(rows_by_contract.values())
    finally:
        workbook.close()


# -- finding the table ---------------------------------------------------------------------

def _normalise_header(value) -> str:
    return re.sub(r"[^a-z0-9]", "", str(value).lower()) if value is not None else ""


def _match_columns(header_cells) -> dict[str, int]:
    """Map our column names to cell positions for one candidate header row."""
    columns = {}
    for position, cell in enumerate(header_cells):
        normalised = _normalise_header(cell)
        for column, aliases in COLUMN_ALIASES.items():
            if normalised in aliases and column not in columns:
                columns[column] = position
    return columns


def _find_contract_table(workbook, file_name: str):
    """Find the (sheet, header row number, column positions) of the contracts table."""
    candidates = []
    headers_seen = {}
    for sheet in workbook.worksheets:
        for row_number, cells in enumerate(
                sheet.iter_rows(max_row=HEADER_SEARCH_ROWS, values_only=True), start=1):
            columns = _match_columns(cells)
            if all(c in columns for c in REQUIRED_COLUMNS):
                candidates.append((sheet, row_number, columns))
                break
            if len(columns) >= 2:  # looks like a header attempt; remember it for the error
                headers_seen[f"{sheet.title}!{row_number}"] = [c for c in cells if c is not None]

    if len(candidates) == 1:
        sheet, row_number, columns = candidates[0]
        log.info("%s: contracts table on sheet '%s', header in row %d", file_name, sheet.title, row_number)
        return candidates[0]
    if not candidates:
        raise WorkbookLayoutError(
            f"{file_name}: no sheet has a header with all of {list(REQUIRED_COLUMNS)}. "
            f"Header-like rows found: {headers_seen or 'none'}. "
            f"If Finance renamed a column, add the new name to COLUMN_ALIASES.")
    raise WorkbookLayoutError(
        f"{file_name}: more than one sheet looks like the contracts table: "
        f"{[c[0].title for c in candidates]}")


# -- parsing rows --------------------------------------------------------------------------

def _parse_row(cells, columns: dict[str, int], sheet_name: str, row_number: int,
               file_name: str) -> SnapshotRow | None:
    def cell(column):
        position = columns.get(column)
        return cells[position] if position is not None and position < len(cells) else None

    contract_id = _clean_text(cell("contract_id"))
    if contract_id is None or not CONTRACT_ID_PATTERN.match(contract_id.upper()):
        # Blank lines, TOTAL rows and footnotes. But a row that has contract dates and no valid
        # id is probably a real contract with a typo: stop rather than silently drop it.
        if cell("start_date") is not None or cell("end_date") is not None:
            raise ValueError(f"{file_name} row {row_number}: has dates but an invalid "
                             f"contract id {cell('contract_id')!r}")
        if contract_id is not None:
            log.info("%s row %d: skipping non-contract row %r", file_name, row_number, contract_id)
        return None

    where = f"{file_name} row {row_number} ({contract_id})"
    terms = ContractTerms(
        customer_id=_required_text(cell("customer_id"), where, "customer id").upper(),
        customer_name=_clean_text(cell("customer_name")),
        start_date=parse_date(cell("start_date"), where),
        end_date=parse_date(cell("end_date"), where),
        commitment_usd=parse_amount(cell("commitment_usd"), where),
        status=_clean_text(cell("status")),
    )
    if terms.end_date < terms.start_date:
        raise ValueError(f"{where}: end date {terms.end_date} is before start date {terms.start_date}")
    return SnapshotRow(contract_id.upper(), terms, sheet_name, row_number)


def _clean_text(value) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def _required_text(value, where: str, what: str) -> str:
    text = _clean_text(value)
    if text is None:
        raise ValueError(f"{where}: missing {what}")
    return text


def parse_date(value, where: str = "") -> date:
    """Accept real Excel dates, Excel serial numbers and a few common text formats."""
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return from_excel(value).date()
    if isinstance(value, str):
        text = value.strip()
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(text, fmt).date()
            except ValueError:
                pass
    raise ValueError(f"{where}: can't read date {value!r}")


def parse_amount(value, where: str = "") -> Decimal:
    """Accept numbers and text such as '$900,000.00' or '900000 USD'. Must be positive."""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        amount = Decimal(str(value))
    elif isinstance(value, str):
        text = re.sub(r"(?i)usd|[$,\s]", "", value)
        try:
            amount = Decimal(text)
        except InvalidOperation:
            raise ValueError(f"{where}: can't read amount {value!r}") from None
    else:
        raise ValueError(f"{where}: can't read amount {value!r}")
    if amount <= 0:
        raise ValueError(f"{where}: commitment must be positive, got {value!r}")
    return amount
